#!/usr/bin/env python3
"""Gera painel_adquirencia.html a partir de 'Report Banco e Adquirencia.xlsx'.

Uso:  python atualizar_painel.py [caminho_do_xlsx] [saida.html]
Requer: pip install openpyxl
Telefone, e-mail e CNPJ dos ECs NAO sao copiados para o HTML.
"""
import sys, json, re, datetime as dt
from pathlib import Path
import openpyxl

AQUI = Path(__file__).resolve().parent
XLSX = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(
    r"C:\Users\GabrielLuizdaSilvaCh\OneDrive - RAFATELLA\BCJ - 08 - Escritório de Processos - "
    r"08 - Escritório de Processos\03 - Indicadores\03 - Status Report CEO\Report BCJ e Adquirencia"
    r"\Report Banco e Adquirencia.xlsx")
SAIDA = Path(sys.argv[2]) if len(sys.argv) > 2 else AQUI / "painel_adquirencia.html"


def val(v):
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return round(v, 6)
    return v


def parse_report(ws):
    # linha 3 = cabecalho do bloco Interno; linha com B='SEÇÃO' apos 'REPORT FRANQUIAS' = Franquias
    header_rows = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == "SEÇÃO"]
    last_col = ws.max_column
    hdr = [ws.cell(header_rows[0], c).value for c in range(5, last_col + 1)]
    cols, cur_month_idx = [], []
    mon = {"JAN": 1, "FEV": 2, "MAR": 3, "ABR": 4, "MAI": 5, "JUN": 6, "JUL": 7, "AGO": 8, "SET": 9, "OUT": 10, "NOV": 11, "DEZ": 12}
    for i, h in enumerate(hdr):
        if isinstance(h, (dt.datetime, dt.date)):
            cols.append({"t": "d", "d": h.strftime("%Y-%m-%d")})
        elif isinstance(h, str) and "SEM" in h:
            cols.append({"t": "w", "l": h})
        elif isinstance(h, str) and re.match(r"^[A-ZÇ]{3}-\d\d$", h):
            cols.append({"t": "m", "l": h, "y": 2000 + int(h[-2:]), "mo": mon[h[:3]]})
        else:
            cols.append({"t": "x", "l": str(h)})
    blocks = {}
    for k, (name, hr) in enumerate(zip(("interna", "franquias"), header_rows[:2])):
        end = (header_rows[k + 1] - 3) if k + 1 < len(header_rows) else ws.max_row
        rows, sec = [], None
        for r in range(hr + 1, end + 1):
            b, c = ws.cell(r, 2).value, ws.cell(r, 3).value
            if b: sec = str(b).strip()
            if not c: continue
            rows.append({"s": sec, "l": str(c), "a": ws.cell(r, 4).value,
                         "v": [val(ws.cell(r, cc).value) for cc in range(5, last_col + 1)]})
        blocks[name] = rows
    return cols, blocks


def table(ws, hdr_row, keep=None):
    heads = [ws.cell(hdr_row, c).value for c in range(1, ws.max_column + 1)]
    out = []
    for r in range(hdr_row + 1, ws.max_row + 1):
        row = [val(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if any(v is not None for v in row): out.append(row)
    heads = [val(h) for h in heads]
    if keep:
        idx = [i for i, h in enumerate(heads) if h is not None and keep(str(h))]
        heads = [heads[i] for i in idx]; out = [[r[i] for i in idx] for r in out]
    return heads, out


wb = openpyxl.load_workbook(XLSX, data_only=True)
cols, blocks = parse_report(wb["Report"])
g = wb["Gestão TPV"]
PII = ("CNPJ", "Responsável", "Telefone", "E-mail")
gh, gr = table(g, 4, keep=lambda h: not h.startswith(PII))
m = re.search(r"atualizado em ([\d/]+ [\d:]+)", str(g["A2"].value or ""))
bh, br = table(wb["Planilha1"], 6) if "Planilha1" in wb.sheetnames else ([], [])
ah, ar = table(wb["Ajustes"], 1)
rg = [[val(c.value) for c in row[:3]] for row in wb["Regras"].iter_rows(min_row=2) if row[0].value]
data = {"updated": m.group(1) if m else None, "source": XLSX.name, "gen": dt.datetime.now().strftime("%d/%m/%Y %H:%M"),
        "gestaoNote": g["A2"].value, "cols": cols, "blocks": blocks,
        "gestao": {"h": gh, "r": gr}, "banco": {"h": bh, "r": br}, "ajustes": {"h": ah, "r": ar}, "regras": rg}
html = (AQUI / "template.html").read_text(encoding="utf-8").replace(
    "/*__DATA__*/null", json.dumps(data, ensure_ascii=False, separators=(",", ":")))
SAIDA.write_text(html, encoding="utf-8")
print(f"OK -> {SAIDA}  (ultima atualizacao da planilha: {data['updated']})")
